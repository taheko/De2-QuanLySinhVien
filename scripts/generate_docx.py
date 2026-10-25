import os
import re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def create_report():
    doc = docx.Document()

    # Set Margins
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Set Base Style Font
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(13)
    normal_style.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    normal_style.paragraph_format.line_spacing = 1.3
    normal_style.paragraph_format.space_after = Pt(6)

    # ------------------ COVER PAGE ------------------
    p_uni = doc.add_paragraph()
    p_uni.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_uni = p_uni.add_run("TRƯỜNG ĐẠI HỌC CÔNG NGHỆ THÔNG TIN VÀ TRUYỀN THÔNG\nKHOA CÔNG NGHỆ THÔNG TIN")
    r_uni.font.name = 'Times New Roman'
    r_uni.font.size = Pt(13)
    r_uni.font.bold = True
    r_uni.font.color.rgb = RGBColor(0x11, 0x18, 0x27)
    p_uni.paragraph_format.space_after = Pt(40)

    p_title_pre = doc.add_paragraph()
    p_title_pre.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title_pre = p_title_pre.add_run("BÁO CÁO HỌC PHẦN\nTRIỂN KHAI & QUẢN TRỊ HỆ THỐNG PHẦN MỀM")
    r_title_pre.font.name = 'Times New Roman'
    r_title_pre.font.size = Pt(16)
    r_title_pre.font.bold = True
    r_title_pre.font.color.rgb = RGBColor(0x1e, 0x3a, 0x8a)
    p_title_pre.paragraph_format.space_after = Pt(25)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_de = p_title.add_run("Đề tài:\n")
    r_de.font.size = Pt(14)
    r_de.font.bold = True
    r_main = p_title.add_run("ĐỀ 2 – HỆ THỐNG QUẢN LÝ SINH VIÊN")
    r_main.font.size = Pt(20)
    r_main.font.bold = True
    r_main.font.color.rgb = RGBColor(0xb9, 0x1c, 0x1c)
    p_title.paragraph_format.space_after = Pt(70)

    # Student info box
    table_info = doc.add_table(rows=4, cols=2)
    table_info.alignment = WD_TABLE_ALIGNMENT.CENTER
    info_data = [
        ("Sinh viên thực hiện :", "Nguyễn Văn A (Cập nhật theo thông tin SV)"),
        ("Mã sinh viên :", "DTC245220060"),
        ("Lớp :", "ATTT K23A / CNTT K23"),
        ("Giảng viên hướng dẫn :", "ThS. Nguyễn Anh Chuyên")
    ]
    for row_idx, (label, val) in enumerate(info_data):
        row = table_info.rows[row_idx]
        c1 = row.cells[0].paragraphs[0]
        c1.paragraph_format.space_after = Pt(4)
        r1 = c1.add_run(label)
        r1.font.name = 'Times New Roman'
        r1.font.bold = True
        r1.font.size = Pt(13)

        c2 = row.cells[1].paragraphs[0]
        c2.paragraph_format.space_after = Pt(4)
        r2 = c2.add_run(val)
        r2.font.name = 'Times New Roman'
        r2.font.size = Pt(13)

    p_foot = doc.add_paragraph()
    p_foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_foot.paragraph_format.space_before = Pt(80)
    r_foot = p_foot.add_run("THÁI NGUYÊN, NĂM 2026")
    r_foot.font.name = 'Times New Roman'
    r_foot.font.bold = True
    r_foot.font.size = Pt(13)

    doc.add_page_break()

    # ------------------ BODY PARSER ------------------
    md_path = "docs/BAO_CAO_DE2_QUAN_LY_SINH_VIEN.md"
    if not os.path.exists(md_path):
        print(f"File {md_path} not found!")
        return

    with open(md_path, "r", encoding="utf-8") as f:
        lines = f.readlines()

    in_cover = True
    in_table = False
    in_code = False
    table_rows = []
    code_lines = []

    for line in lines:
        raw_line = line.rstrip("\r\n")

        # Skip cover page in markdown body
        if "## LỜI NÓI ĐẦU" in raw_line:
            in_cover = False

        if in_cover:
            continue

        # Handle code blocks
        if raw_line.startswith("```"):
            if not in_code:
                in_code = True
                code_lines = []
            else:
                in_code = False
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(6)
                p.paragraph_format.space_after = Pt(6)
                p.paragraph_format.left_indent = Inches(0.2)
                r = p.add_run("\n".join(code_lines))
                r.font.name = "Consolas"
                r.font.size = Pt(10)
                r.font.color.rgb = RGBColor(0x1e, 0x29, 0x3b)
            continue

        if in_code:
            code_lines.append(raw_line)
            continue

        # Handle tables
        if raw_line.startswith("|") and raw_line.endswith("|"):
            # Check separator row
            if re.match(r"^\|(\s*:?-+:?\s*\|)+$", raw_line):
                continue
            cols = [c.strip() for c in raw_line.strip("|").split("|")]
            table_rows.append(cols)
            in_table = True
            continue
        elif in_table:
            # End of table
            if table_rows:
                num_cols = max(len(r) for r in table_rows)
                t = doc.add_table(rows=len(table_rows), cols=num_cols)
                t.alignment = WD_TABLE_ALIGNMENT.CENTER
                for r_idx, row in enumerate(table_rows):
                    is_header = (r_idx == 0)
                    for c_idx, val in enumerate(row):
                        if c_idx < num_cols:
                            cell = t.cell(r_idx, c_idx)
                            cp = cell.paragraphs[0]
                            cp.paragraph_format.space_after = Pt(2)
                            cp.paragraph_format.space_before = Pt(2)
                            r = cp.add_run(val)
                            r.font.name = 'Times New Roman'
                            r.font.size = Pt(11)
                            if is_header:
                                r.font.bold = True
                                set_cell_background(cell, "F1F5F9")
                            set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
                doc.add_paragraph().paragraph_format.space_after = Pt(4)
            table_rows = []
            in_table = False

        # Headings
        if raw_line.startswith("## "):
            h = doc.add_heading(level=1)
            h.paragraph_format.space_before = Pt(16)
            h.paragraph_format.space_after = Pt(6)
            r = h.add_run(raw_line[3:].strip())
            r.font.name = 'Times New Roman'
            r.font.size = Pt(16)
            r.font.bold = True
            r.font.color.rgb = RGBColor(0x1e, 0x3a, 0x8a)
        elif raw_line.startswith("### "):
            h = doc.add_heading(level=2)
            h.paragraph_format.space_before = Pt(12)
            h.paragraph_format.space_after = Pt(4)
            r = h.add_run(raw_line[4:].strip())
            r.font.name = 'Times New Roman'
            r.font.size = Pt(14)
            r.font.bold = True
            r.font.color.rgb = RGBColor(0x0f, 0x17, 0x2a)
        elif raw_line.startswith("#### "):
            h = doc.add_heading(level=3)
            h.paragraph_format.space_before = Pt(8)
            h.paragraph_format.space_after = Pt(2)
            r = h.add_run(raw_line[5:].strip())
            r.font.name = 'Times New Roman'
            r.font.size = Pt(13)
            r.font.bold = True
            r.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
        elif raw_line.strip() == "":
            continue
        elif raw_line.startswith("- ") or raw_line.startswith("* "):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.25
            r = p.add_run(raw_line[2:].strip())
            r.font.name = 'Times New Roman'
            r.font.size = Pt(13)
        else:
            p = doc.add_paragraph()
            p.paragraph_format.line_spacing = 1.3
            p.paragraph_format.space_after = Pt(6)
            # Simple bold parsing **text**
            parts = re.split(r'(\*\*.*?\*\*)', raw_line)
            for part in parts:
                if part.startswith("**") and part.endswith("**"):
                    r = p.add_run(part[2:-2])
                    r.font.bold = True
                else:
                    r = p.add_run(part)
                r.font.name = 'Times New Roman'
                r.font.size = Pt(13)

    out_docx = "Bao_cao_De2_Quan_Ly_Sinh_Vien.docx"
    doc.save(out_docx)
    print(f"Generated {out_docx} successfully! Size: {os.path.getsize(out_docx)} bytes")

if __name__ == "__main__":
    create_report()
